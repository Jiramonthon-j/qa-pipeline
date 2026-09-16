#!/usr/bin/env python3
"""
ป้องกันปัญหาการแบ่งหน้าที่อ่านยากในไฟล์ .docx ที่ pandoc สร้างไว้:

1. หัวข้อ (Heading) ไม่ตกไปอยู่โดดเดี่ยวท้ายหน้าโดยที่เนื้อหาข้างล่างเริ่มหน้าถัดไป
   (keep-with-next) — ตั้งให้ทุกย่อหน้าหัวข้อ "ติด" กับย่อหน้าถัดไปเสมอ
2. เปิด widow/orphan control ให้ย่อหน้าเนื้อหาทั่วไป กันบรรทัดเดียวของย่อหน้าหลุดไปอยู่
   คนละหน้ากับส่วนที่เหลือ

สคริปต์นี้แก้ไฟล์ .docx ที่ pandoc export ไว้แล้วเท่านั้น ไม่ได้บังคับให้ทั้ง section
อยู่หน้าเดียวกันเสมอ (section ยาวๆ เช่น "สรุปจากแต่ละแหล่ง" ยังขึ้นหน้าใหม่ได้ตามปกติ
ถ้ายาวเกิน 1 หน้า) — จุดประสงค์คือกันแค่ "หัวข้อลอยเดี่ยวท้ายหน้า" ซึ่งเป็นจุดที่อ่านสับสนที่สุด

ใช้งาน:
    python3 fix_page_breaks.py "<path .docx>"

แก้ไฟล์ในที่เดิม (in-place) — ควรรันหลัง pandoc export (และหลัง add_table_borders.py
ถ้ามี) เสร็จแล้วเท่านั้น ห้ามรันกับไฟล์ .docx ที่เป็นไฟล์หลัก (source of truth ยังคงเป็น
.md เท่านั้น ไฟล์นี้แก้แค่สำเนา .docx เพื่อความสวยงาม)
"""
import sys

from docx import Document


def main():
    if len(sys.argv) != 2:
        print("usage: python3 fix_page_breaks.py <path .docx>")
        sys.exit(1)

    path = sys.argv[1]
    doc = Document(path)

    heading_count = 0
    body_count = 0
    for para in doc.paragraphs:
        style_name = para.style.name if para.style else ""
        pf = para.paragraph_format
        if style_name.startswith("Heading") or style_name == "Title":
            pf.keep_with_next = True
            heading_count += 1
        else:
            pf.widow_control = True
            body_count += 1

    doc.save(path)
    print(
        f"ตั้ง keep-with-next ให้ {heading_count} หัวข้อ และเปิด widow/orphan control "
        f"ให้ {body_count} ย่อหน้าใน {path} เรียบร้อย"
    )


if __name__ == "__main__":
    main()
