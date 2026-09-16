#!/usr/bin/env python3
"""
read_workbook.py — สคริปต์อ่านอย่างเดียว (read-only) สำหรับ Skill 07 (result-analysis)

ตั้งใจให้เป็นสคริปต์แยกจาก qa_workbook.py (ที่ใช้ใน Skill 03/04/05/06a) เพราะ Skill 07
เป็น Skill อ่าน-สรุปผลเท่านั้น ไม่ควรมีความสามารถแก้ไข Workbook ปนอยู่ในสคริปต์เดียวกันเลย
(ป้องกัน Skill นี้เผลอเรียกคำสั่งแก้ไขโดยไม่ตั้งใจ)

คำสั่งที่รองรับ:
  read-cases --path <workbook.xlsx>
      คืนค่า Test Case ทุกแถวเป็น JSON list

  read-sheet --path <workbook.xlsx> --sheet "<ชื่อชีต>"
      คืนค่าทุกแถวของชีตที่ระบุ (เช่น "Env & Config", "Priority & Risk Matrix",
      "Requirement Matrix") เป็น JSON list of list — ใช้สำหรับอ่านชีตอื่นที่ไม่ใช่ Test Case
"""
import argparse
import json

import openpyxl

TC_COLUMNS = [
    "Test Case ID", "Module", "Feature", "Test Type", "Priority",
    "Test Scenario", "Pre-condition", "Test Step", "Test Data",
    "Expected Result", "Actual Result", "Status", "Issue link",
    "Test By", "Execution Date", "Remarks", "Test Photo",
]


def read_cases(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Test Case"]
    cases = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        cases.append({col: row[i] for i, col in enumerate(TC_COLUMNS)})
    return cases


def read_sheet(path, sheet_name):
    wb = openpyxl.load_workbook(path, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"ไม่พบชีตชื่อ '{sheet_name}' ในไฟล์นี้ — ชีตที่มี: {wb.sheetnames}")
    ws = wb[sheet_name]
    rows = []
    for row in ws.iter_rows(values_only=True):
        if any(c is not None for c in row):
            rows.append(list(row))
    return rows


def main():
    p = argparse.ArgumentParser()
    sub = p.add_subparsers(dest="cmd", required=True)

    p_cases = sub.add_parser("read-cases")
    p_cases.add_argument("--path", required=True)

    p_sheet = sub.add_parser("read-sheet")
    p_sheet.add_argument("--path", required=True)
    p_sheet.add_argument("--sheet", required=True)

    args = p.parse_args()

    if args.cmd == "read-cases":
        print(json.dumps(read_cases(args.path), ensure_ascii=False, indent=2, default=str))
    elif args.cmd == "read-sheet":
        print(json.dumps(read_sheet(args.path, args.sheet), ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
