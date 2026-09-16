#!/usr/bin/env python3
"""
build_final_report.py — สคริปต์สร้างรายงานฉบับสุดท้ายจริง สำหรับ Skill 08 (qa-report-generator)

สคริปต์นี้ทำหน้าที่ "render" เท่านั้น — การตัดสินใจ Go/No-Go เป็นวิจารณญาณ/กฎที่ระบุไว้ใน
ขั้นตอนที่ 3 ของ SKILL.md ไม่ใช่ logic ในสคริปต์นี้ (สคริปต์แค่วาดผลลัพธ์ที่ตัดสินใจมาแล้ว)

คำสั่ง:
  build --data <report.json> --out <08-qa-report.docx>

data schema:
{
  "feature": str, "report_date": str, "deadline": str, "environment": str,
  "verdict": "Go" | "Conditional Go" | "No-Go", "is_provisional": bool,
  "provisional_reason": str,  # จำเป็นถ้า is_provisional=true — ระบุสาเหตุให้ครบทุกข้อ (RTM ไม่เสร็จ
                               # และ/หรือ Stage 07 ยังเป็น Partial ตาม Gate ข้อ 3a/3b ของ SKILL.md ขั้นตอนที่ 1)
                               # คั่นหลายสาเหตุด้วย "; " ห้ามใช้ข้อความ hardcode ตายตัวแบบเดิมที่อ้างแค่ RTM
  "verdict_reason": str,
  "coverage": {"total": int, "executed": int, "coverage_pct": float},
  "status_priority_matrix": {status: {"P0":int,"P1":int,"P2":int,"P3":int}},
  "known_issues": [{"category": str, "tc": str|None, "priority": str|None, "detail": str, "recommendation": str}],
  "browsers_tested": [str], "browsers_required": [str], "browser_complete": bool,
  "days_vs_deadline": str, "deadline_passed": bool,
  "next_steps": [str],
}
"""
import argparse
import json

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn

VERDICT_COLORS = {
    "Go": RGBColor(0x00, 0x80, 0x00),
    "Conditional Go": RGBColor(0xCC, 0x66, 0x00),
    "No-Go": RGBColor(0xC0, 0x00, 0x00),
}

# ความกว้างคอลัมน์ของตาราง Known Issues (หมวด | Test Case | Priority | รายละเอียด | คำแนะนำ)
# ยืนยันแล้วผ่าน Mockup — ห้ามปล่อยให้ Word/LibreOffice หารความกว้างเท่ากันทั้ง 5 คอลัมน์แบบเดิม
# เพราะจะทำให้คอลัมน์รายละเอียด/คำแนะนำแคบเกินไปจนข้อความแตกเป็นบรรทัดสั้นๆ ทีละ 1-2 คำ อ่านยาก
KNOWN_ISSUES_COL_WIDTHS = [Inches(1.1), Inches(0.65), Inches(0.45), Inches(2.25), Inches(2.05)]


def _set_cell_width(cell, width):
    cell.width = width
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn('w:tcW'))
    if tcW is None:
        tcW = tcPr.makeelement(qn('w:tcW'), {})
        tcPr.append(tcW)
    tcW.set(qn('w:w'), str(width.twips))
    tcW.set(qn('w:type'), 'dxa')


def _set_col_widths(table, widths):
    """ตั้งความกว้างคอลัมน์แบบตายตัว — จุดสำคัญที่พลาดง่ายของ python-docx: ตั้งแค่ cell.width
    ทีละ cell ไม่พอ เพราะความกว้างที่ Word/LibreOffice ใช้จริงตอน layout="fixed" มาจาก w:tblGrid/
    w:gridCol ของทั้งตาราง ไม่ใช่ w:tcW ของแต่ละ cell — ต้องตั้ง table.columns[i].width ด้วยเสมอ
    (นอกเหนือจาก cell.width) ไม่งั้นคอลัมน์จะยังถูกหารความกว้างเท่าๆ กันเหมือนเดิมทุกคอลัมน์"""
    table.autofit = False
    table.allow_autofit = False
    for i, w in enumerate(widths):
        table.columns[i].width = w
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            _set_cell_width(cell, w)


def build_report(data, out_path):
    doc = Document()

    doc.add_heading(f"QA Final Report — {data['feature']}", level=0)
    meta = doc.add_paragraph()
    meta.add_run(f"วันที่ออกรายงาน: {data['report_date']}\n")
    meta.add_run(f"Deadline: {data['deadline']}\n")
    meta.add_run(f"Environment: {data.get('environment', '-')}")

    doc.add_heading("ผลสรุป", level=1)
    verdict_para = doc.add_paragraph()
    verdict_run = verdict_para.add_run(data["verdict"])
    verdict_run.bold = True
    verdict_run.font.size = Pt(20)
    verdict_run.font.color.rgb = VERDICT_COLORS.get(data["verdict"], RGBColor(0, 0, 0))
    if data.get("is_provisional"):
        # แยกเป็นย่อหน้าของตัวเอง ขนาดตัวอักษรปกติ — เดิมเคยต่อท้าย verdict_run เดียวกันที่ font 20pt
        # ทำให้สาเหตุ Provisional (ซึ่งอาจยาว โดยเฉพาะเมื่อมีหลายสาเหตุพร้อมกันตาม Gate 3a/3b) กลาย
        # เป็นตัวอักษรยักษ์ทั้งท่อนไปด้วย อ่านยากและดูไม่เป็นมืออาชีพ — และเดิม hardcode ข้อความว่า
        # "RTM ยังไม่เสร็จสมบูรณ์" เสมอไม่ว่าสาเหตุจริงจะเป็นอะไร ทำให้เคสที่ Provisional เพราะ Stage 07
        # ยังเป็น Partial (ไม่ใช่เพราะ RTM) แสดงข้อความผิดจากความจริง — แก้ทั้งสองจุดพร้อมกัน: รับสาเหตุ
        # จริงจาก provisional_reason (ดู schema ด้านบน) และวางเป็นย่อหน้าแยกขนาดพอดี
        reason = data.get("provisional_reason") or "ยังไม่ผ่าน Gate ครบทุกข้อ"
        prov_para = doc.add_paragraph()
        prov_run = prov_para.add_run(f"⚠ ชั่วคราว (Provisional): {reason}")
        prov_run.bold = True
        prov_run.font.size = Pt(12)
        prov_run.font.color.rgb = RGBColor(0xCC, 0x66, 0x00)
    doc.add_paragraph(f"เหตุผล: {data['verdict_reason']}")

    doc.add_heading("Test Coverage", level=1)
    cov = data["coverage"]
    doc.add_paragraph(
        f"ทดสอบไปแล้ว {cov['executed']} จาก {cov['total']} เคส ({cov['coverage_pct']:.1f}%)"
    )
    spm = data["status_priority_matrix"]
    statuses = list(spm.keys())
    table = doc.add_table(rows=1 + len(statuses), cols=5)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, h in enumerate(["Status", "P0", "P1", "P2", "P3"]):
        hdr[i].text = h
    for r, status in enumerate(statuses, start=1):
        cells = table.rows[r].cells
        cells[0].text = status
        for c, prio in enumerate(["P0", "P1", "P2", "P3"], start=1):
            cells[c].text = str(spm[status].get(prio, 0))

    doc.add_heading("Known Issues / ความเสี่ยงที่ยังค้างอยู่", level=1)
    issues = data.get("known_issues", [])
    if not issues:
        doc.add_paragraph("ไม่มี Known Issue ค้างอยู่ในรอบนี้")
    else:
        t2 = doc.add_table(rows=1 + len(issues), cols=5)
        t2.style = "Light Grid Accent 1"
        h2 = t2.rows[0].cells
        for i, h in enumerate(["หมวด", "Test Case", "Priority", "รายละเอียด", "คำแนะนำ"]):
            h2[i].text = h
            for p in h2[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
        for r, issue in enumerate(issues, start=1):
            cells = t2.rows[r].cells
            cell_values = [
                issue.get("category", ""),
                issue.get("tc") or "-",
                issue.get("priority") or "-",
                issue.get("detail", ""),
                issue.get("recommendation", ""),
            ]
            for c, val in enumerate(cell_values):
                cells[c].text = val
                cells[c].vertical_alignment = WD_ALIGN_VERTICAL.TOP
                for p in cells[c].paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(10)
        _set_col_widths(t2, KNOWN_ISSUES_COL_WIDTHS)

    doc.add_heading("Environment / เวอร์ชันที่ทดสอบ", level=1)
    doc.add_paragraph(f"Browser ที่ทดสอบแล้ว: {', '.join(data.get('browsers_tested', [])) or '-'}")
    doc.add_paragraph(f"Browser ที่ต้องทดสอบตาม Env & Config: {', '.join(data.get('browsers_required', [])) or '-'}")
    complete_para = doc.add_paragraph()
    complete_run = complete_para.add_run(
        "ทดสอบครบทุก Browser ที่กำหนดแล้ว" if data.get("browser_complete") else "⚠ ยังทดสอบไม่ครบทุก Browser ที่กำหนด"
    )
    if not data.get("browser_complete"):
        complete_run.bold = True
        complete_run.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)

    doc.add_heading("วันที่เสร็จจริงเทียบ Deadline", level=1)
    deadline_para = doc.add_paragraph()
    deadline_run = deadline_para.add_run(data.get("days_vs_deadline", "-"))
    if data.get("deadline_passed"):
        deadline_run.bold = True
        deadline_run.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)
        warn_para = doc.add_paragraph()
        warn_run = warn_para.add_run("⚠ เลย Deadline ที่กำหนดไว้แล้ว — โปรดพิจารณาประกอบการตัดสินใจ")
        warn_run.bold = True
        warn_run.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)

    doc.add_heading("Next Steps", level=1)
    for step in data.get("next_steps", []):
        doc.add_paragraph(step, style="List Bullet")

    doc.add_heading("ช่องเซ็นอนุมัติ", level=1)
    approval = doc.add_table(rows=3, cols=4)
    approval.style = "Table Grid"
    ah = approval.rows[0].cells
    for i, h in enumerate(["Role", "ชื่อ", "ลายเซ็น", "วันที่"]):
        ah[i].text = h
    for role, row_idx in [("QA Lead", 1), ("PM / Product Owner", 2)]:
        approval.rows[row_idx].cells[0].text = role
        for c in range(1, 4):
            approval.rows[row_idx].cells[c].text = ""

    doc.save(out_path)
    return out_path


def main():
    p = argparse.ArgumentParser()
    sub = p.add_subparsers(dest="cmd", required=True)
    p_build = sub.add_parser("build")
    p_build.add_argument("--data", required=True)
    p_build.add_argument("--out", required=True)
    args = p.parse_args()

    with open(args.data, encoding="utf-8") as f:
        data = json.load(f)

    out = build_report(data, args.out)
    print(json.dumps({"written": out}, ensure_ascii=False))


if __name__ == "__main__":
    main()
