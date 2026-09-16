#!/usr/bin/env python3
"""
จัดความกว้างคอลัมน์ตารางใน .docx ที่ pandoc สร้างไว้ให้อ่านง่ายขึ้น

ปัญหาเดิม (v1): pandoc หารความกว้างคอลัมน์เท่า ๆ กันทุกคอลัมน์เสมอ ทำให้คอลัมน์ข้อความยาว
(เช่น "รายละเอียด", "ข้อเสนอแนะ/Mitigation") แคบเกินไปจนคำถูกตัดกลางคำ

ปัญหาที่พบเพิ่มเติม (v2 — จากตาราง Risks 7 คอลัมน์ของ requirement-review/test-plan): แม้จะปรับ
ความกว้างตามเนื้อหาแล้ว คอลัมน์สั้นที่มีคำภาษาอังกฤษคำเดียวยาวๆ (เช่น "Likelihood", "Medium",
"Operational") ยังโดนบีบจนต่ำกว่าความกว้างที่คำนั้นต้องการ ทำให้ Word ตัดคำกลางคำแบบไม่มีจุดเชื่อม
("Likeliho" ขึ้นบรรทัดใหม่เป็น "od") ซึ่งอ่านยากกว่าตารางที่ไม่ได้ปรับความกว้างเลยด้วยซ้ำ — ต่างจาก
ข้อความไทยที่ตัดขึ้นบรรทัดใหม่ตรงไหนก็ยังอ่านได้ปกติ (ภาษาไทยไม่มีช่องว่างระหว่างคำอยู่แล้ว)

วิธีแก้ (v2): เพิ่ม 2 กลไกเข้ามาคู่กัน
1. หาความกว้างขั้นต่ำต่อคอลัมน์แบบ "รู้เนื้อหา" — สแกนหาคำ/โทเคนภาษาอังกฤษ-ตัวเลขที่ยาวที่สุดใน
   แต่ละคอลัมน์ (ทั้ง header และข้อมูล) แล้วคำนวณความกว้างขั้นต่ำที่พอให้คำนั้นอยู่บรรทัดเดียวได้
   โดยไม่ไปแตะโทเคนภาษาไทย (เพราะภาษาไทยตัดบรรทัดที่ไหนก็ได้อยู่แล้ว ไม่ต้องกันความกว้างให้)
2. ลดขนาดฟอนต์ในตารางลงเหลือ 10pt (จากค่า default ของ Word/LibreOffice ที่มักใหญ่กว่านี้) เพื่อเพิ่ม
   ความหนาแน่นตัวอักษรต่อนิ้ว ช่วยให้ตารางที่มีหลายคอลัมน์ (เช่น Risks 7 คอลัมน์) มีที่พอสำหรับทุกคอลัมน์
   โดยไม่ต้องบีบคอลัมน์ข้อความยาวจนอ่านไม่ออก — ยังคงอ่านง่ายกว่าตัวอักษรขนาดปกติที่ต้องตัดคำกลางคำ

ทั้งสองกลไกนี้เป็น heuristic ทั่วไปที่คำนวณจากเนื้อหาจริง ไม่ได้ hardcode ตามชื่อคอลัมน์หรือ
โครงสร้างตารางเฉพาะของ Skill ไหน จึงใช้ได้กับทุกตารางที่มีอยู่ในไฟล์ (Skill นี้มีหลายตารางโครงสร้าง
ต่างกัน)

ใช้งาน:
    python3 resize_table_columns.py "<path .docx>"

แก้ไฟล์ .docx ในที่เดิม (in-place) — ควรรันหลัง pandoc export และก่อน add_table_borders.py
"""
import re
import sys

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn

USABLE_WIDTH_IN = 6.5  # Letter, margin 1" ทั้งสองข้าง — ตั้ง section ให้ตรงกันเสมอ (ดู _ensure_page_setup)
MIN_COL_WIDTH_IN = 0.55  # floor เผื่อคอลัมน์ที่ไม่มีโทเคนภาษาอังกฤษยาวๆ เลย (เช่นตัวเลขล้วน/ว่าง)
MAX_COL_WIDTH_IN = 3.2
TABLE_FONT_PT = 10  # เล็กกว่าฟอนต์เนื้อหาปกติเล็กน้อย เพิ่มความหนาแน่นให้ตารางหลายคอลัมน์พอมีที่ทุกคอลัมน์

# โทเคนภาษาอังกฤษ/ตัวเลข/ขีดกลาง ต่อเนื่องกัน — ใช้หาคำที่ "ตัดกลางคำไม่ได้อย่างสวยงาม" ถ้าคอลัมน์แคบไป
_LATIN_TOKEN_RE = re.compile(r"[A-Za-z0-9_][A-Za-z0-9_\-]*")
_CHAR_WIDTH_IN = 0.072  # ความกว้างเฉลี่ยโดยประมาณต่อตัวอักษรที่ TABLE_FONT_PT (ปรับตามผลจริงถ้าจำเป็น)
_CELL_PADDING_IN = 0.18  # ระยะขอบซ้าย-ขวาภายใน cell รวมกัน (ค่า default ของ Word)


def _ensure_page_setup(doc):
    """pandoc ไม่ได้ระบุ page size/margin ไว้ในไฟล์เสมอไป (ปล่อยให้ผู้เปิดใช้ default ของตัวเอง)
    ทำให้คำนวณความกว้างคอลัมน์จากสมมติฐาน 6.5" คลาดเคลื่อนได้ถ้าเปิดคนละโปรแกรม — ล็อกให้ชัดเจนไปเลย
    เป็น Letter portrait margin 1" ทุกด้าน เพื่อให้ความกว้างคอลัมน์ที่คำนวณไว้ตรงกับพื้นที่จริงเสมอ"""
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)


def _cell_text(cell):
    return " ".join(p.text for p in cell.paragraphs).strip()


def _col_weights(table):
    ncols = len(table.columns)
    header_lens = [len(_cell_text(c)) for c in table.rows[0].cells[:ncols]]
    data_lens_by_col = [[] for _ in range(ncols)]
    for row in table.rows[1:]:
        for i, cell in enumerate(row.cells[:ncols]):
            data_lens_by_col[i].append(len(_cell_text(cell)))
    weights = []
    for i in range(ncols):
        avg_data = sum(data_lens_by_col[i]) / len(data_lens_by_col[i]) if data_lens_by_col[i] else 0
        weights.append(max(header_lens[i], avg_data, 1))
    return weights


def _longest_latin_token_len(text):
    tokens = _LATIN_TOKEN_RE.findall(text)
    return max((len(t) for t in tokens), default=0)


def _col_min_widths(table):
    """ความกว้างขั้นต่ำต่อคอลัมน์ — กันไม่ให้คำภาษาอังกฤษ/ตัวเลขคำเดียวถูกตัดกลางคำ
    (ดูคำอธิบายปัญหาที่หัวไฟล์) ไม่นับความยาวข้อความภาษาไทยเพราะตัดบรรทัดที่ไหนก็อ่านได้ปกติ"""
    ncols = len(table.columns)
    mins = [MIN_COL_WIDTH_IN] * ncols
    for row in table.rows:
        for i, cell in enumerate(row.cells[:ncols]):
            tok_len = _longest_latin_token_len(_cell_text(cell))
            if tok_len:
                needed = tok_len * _CHAR_WIDTH_IN + _CELL_PADDING_IN
                if needed > mins[i]:
                    mins[i] = min(needed, MAX_COL_WIDTH_IN)
    return mins


def _water_fill(weights, total, min_widths, max_w):
    n = len(weights)
    widths = list(min_widths)
    remaining = total - sum(min_widths)
    if remaining <= 0:
        # ผลรวมความกว้างขั้นต่ำที่ต้องการเกินความกว้างหน้ากระดาษไปแล้ว (ตารางคอลัมน์เยอะมากจริงๆ)
        # ย่อทุกคอลัมน์ลงตามสัดส่วนแทนที่จะปล่อยให้ตารางล้นหน้ากระดาษ
        total_min = sum(min_widths)
        if total_min <= 0:
            return [total / n] * n
        scale = total / total_min
        return [w * scale for w in min_widths]
    active = set(i for i in range(n) if widths[i] < max_w)
    while remaining > 1e-6 and active:
        total_w = sum(weights[i] for i in active)
        if total_w <= 0:
            share = remaining / len(active)
            for i in active:
                widths[i] += share
            remaining = 0
            break
        capped = []
        alloc = {}
        for i in active:
            add = remaining * (weights[i] / total_w)
            if widths[i] + add >= max_w:
                alloc[i] = max_w - widths[i]
                capped.append(i)
            else:
                alloc[i] = add
        for i in active:
            widths[i] += alloc[i]
        remaining -= sum(alloc.values())
        if not capped:
            break
        active -= set(capped)
    return widths


def _set_cell_width(cell, width):
    cell.width = width
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = tcPr.makeelement(qn("w:tcW"), {})
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(width.twips))
    tcW.set(qn("w:type"), "dxa")


def _apply_widths(table, widths_in):
    # จุดสำคัญของ python-docx: ต้องตั้งทั้ง table.columns[i].width (กำหนด w:tblGrid) และ
    # cell.width ของทุก cell (กำหนด w:tcW) คู่กันเสมอ — ตั้งแค่อย่างใดอย่างหนึ่งจะไม่มีผลจริง
    table.autofit = False
    table.allow_autofit = False
    widths = [Inches(w) for w in widths_in]
    for i, w in enumerate(widths):
        if i < len(table.columns):
            table.columns[i].width = w
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            _set_cell_width(cell, w)


def _set_table_font_size(table, pt):
    size = Pt(pt)
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = size


def resize_tables(doc):
    n = 0
    for table in doc.tables:
        if not table.rows or not table.columns:
            continue
        _set_table_font_size(table, TABLE_FONT_PT)
        weights = _col_weights(table)
        min_widths = _col_min_widths(table)
        widths_in = _water_fill(weights, USABLE_WIDTH_IN, min_widths, MAX_COL_WIDTH_IN)
        _apply_widths(table, widths_in)
        n += 1
    return n


def main():
    if len(sys.argv) != 2:
        print("usage: python3 resize_table_columns.py <path .docx>", file=sys.stderr)
        sys.exit(1)
    path = sys.argv[1]
    doc = Document(path)
    _ensure_page_setup(doc)
    n = resize_tables(doc)
    doc.save(path)
    print(f"ปรับความกว้างคอลัมน์ + ขนาดฟอนต์ให้ {n} ตารางใน {path} เรียบร้อย")


if __name__ == "__main__":
    main()
