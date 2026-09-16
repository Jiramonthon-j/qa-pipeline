#!/usr/bin/env python3
"""
build_docx_reports.py — สคริปต์สร้างไฟล์ Word จริงสำหรับ Skill 07 (result-analysis)

สคริปต์นี้ทำหน้าที่ "render" เท่านั้น — การจัดกลุ่ม/วิเคราะห์ Root Cause เป็นวิจารณญาณของ
ผู้ที่รัน Skill (Claude/Gemini) ตามขั้นตอนที่ 3 ใน SKILL.md ไม่ใช่ logic ในสคริปต์นี้

คำสั่งที่รองรับ:
  summary --data <summary.json> --out <07-result-analysis.docx>
  photo-evidence --data <photo_evidence.json> --out <07-photo-evidence.docx>

ดู docstring ของแต่ละฟังก์ชันสำหรับ schema ของไฟล์ JSON ที่ต้องเตรียม
"""
import argparse
import json
import os

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

STATUS_COLORS = {
    "Pass": RGBColor(0x00, 0x80, 0x00),
    "Fail": RGBColor(0xC0, 0x00, 0x00),
    "Blocked": RGBColor(0xCC, 0x66, 0x00),
    "Rejected / Not a Bug": RGBColor(0x80, 0x80, 0x80),
    "not start": RGBColor(0x80, 0x80, 0x80),
}


def build_summary(data, out_path):
    """
    data schema:
    {
      "feature": str, "analysis_date": str, "environment": str,
      "browsers_tested": [str], "is_partial": bool, "partial_note": str|None,
      "coverage": {"total": int, "executed": int, "not_start": int, "coverage_pct": float},
      "status_priority_matrix": {status: {"P0":int,"P1":int,"P2":int,"P3":int}},
      "root_causes": [{"title":str,"category":str,"affected_tc":[str],"max_priority":str,"description":str}],
      "must_fix_before_close": [{"tc":str,"priority":str,"summary":str}]
    }
    """
    doc = Document()

    doc.add_heading(f"Result Analysis — {data['feature']}", level=0)
    p = doc.add_paragraph()
    p.add_run(f"วันที่วิเคราะห์: {data['analysis_date']}\n").bold = False
    p.add_run(f"Environment: {data.get('environment', '-')}\n")
    p.add_run(f"Browser ที่ทดสอบแล้ว: {', '.join(data.get('browsers_tested', [])) or '-'}\n")

    if data.get("is_partial"):
        warn = doc.add_paragraph()
        run = warn.add_run(f"⚠ การวิเคราะห์นี้เป็นแบบ PARTIAL: {data.get('partial_note', '')}")
        run.bold = True
        run.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)

    doc.add_heading("สรุปภาพรวม", level=1)
    cov = data["coverage"]
    doc.add_paragraph(
        f"ทดสอบไปแล้ว {cov['executed']} จาก {cov['total']} เคส "
        f"({cov['coverage_pct']:.1f}%) — ยังไม่ได้ทดสอบ {cov['not_start']} เคส"
    )

    doc.add_heading("ตาราง Status × Priority", level=2)
    spm = data["status_priority_matrix"]
    statuses = list(spm.keys())
    table = doc.add_table(rows=1 + len(statuses), cols=5)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, h in enumerate(["Status", "P0", "P1", "P2", "P3"]):
        hdr[i].text = h
    for r, status in enumerate(statuses, start=1):
        cells = table.rows[r].cells
        cells[0].text = status
        for c, prio in enumerate(["P0", "P1", "P2", "P3"], start=1):
            cells[c].text = str(spm[status].get(prio, 0))

    doc.add_heading("ปัญหาหลักที่พบ (Root Cause)", level=1)
    if not data.get("root_causes"):
        doc.add_paragraph("ไม่พบปัญหาใดๆ ในรอบทดสอบนี้")
    for rc in data.get("root_causes", []):
        doc.add_heading(rc["title"], level=2)
        meta = doc.add_paragraph()
        meta.add_run(f"หมวด: {rc['category']}  |  Priority สูงสุดที่กระทบ: {rc['max_priority']}\n").italic = True
        meta.add_run(f"กระทบ Test Case: {', '.join(rc['affected_tc'])}")
        doc.add_paragraph(rc["description"])

    doc.add_heading("ประเด็นที่ต้องแก้ก่อนปิดจบ (P0/P1 ที่ยัง Fail จริง)", level=1)
    mfb = data.get("must_fix_before_close", [])
    if not mfb:
        doc.add_paragraph("ไม่มีประเด็นระดับ P0/P1 ที่ยัง Fail ค้างอยู่ในรอบนี้")
    else:
        table2 = doc.add_table(rows=1 + len(mfb), cols=3)
        table2.style = "Light Grid Accent 1"
        h2 = table2.rows[0].cells
        h2[0].text, h2[1].text, h2[2].text = "Test Case ID", "Priority", "สรุปปัญหา"
        for r, item in enumerate(mfb, start=1):
            cells = table2.rows[r].cells
            cells[0].text = item["tc"]
            cells[1].text = item["priority"]
            cells[2].text = item["summary"]

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.add_run("ขั้นตอนถัดไป: ไปต่อ qa-reconcile (RTM) เพื่อตรวจสอบ Traceability ก่อนสรุปรายงานฉบับสุดท้าย").italic = True

    doc.save(out_path)
    return out_path


def _add_issue_points(doc, issue_point):
    """เขียนย่อหน้า 'จุดที่ผิด:' แล้วตามด้วยเนื้อหา — รองรับทั้ง str เดิม (backward compatible)
    และ list ของหลายจุด (ใหม่) เพื่อไม่ให้หลายประเด็นไปติดกันในบรรทัดเดียวจนอ่านยาก

    - str/None: ขึ้นบรรทัดเดียวตามเดิม ไม่มี bullet (กรณีมีจุดผิดแค่ 1 จุด หรือ '-'/'ยังไม่ได้ทดสอบ')
    - list ที่ไม่ว่าง: แต่ละจุดขึ้นบรรทัดของตัวเอง นำหน้าด้วย '•' และเยื้องซ้ายเล็กน้อย
    - list ว่าง: แสดง '-'
    """
    label = doc.add_paragraph()
    label.add_run("จุดที่ผิด:").bold = True

    if isinstance(issue_point, list):
        if not issue_point:
            doc.add_paragraph("-")
            return
        for point in issue_point:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.add_run(f"•  {point}")
    else:
        p = doc.add_paragraph(str(issue_point) if issue_point else "-")
        p.paragraph_format.left_indent = Inches(0.25)


def build_photo_evidence(data, out_path):
    """
    data schema:
    {
      "feature": str,
      "cases": [
        {
          "tc_id": str, "scenario": str, "priority": str, "status": str,
          "description": str,
          "issue_point": str | list[str],
          # str: จุดผิดจุดเดียว (หรือ "-" / "ยังไม่ได้ทดสอบ") — ขึ้นบรรทัดเดียวตามเดิม
          # list[str]: มีหลายจุด — แต่ละจุดจะถูกแยกขึ้นบรรทัดของตัวเองพร้อม bullet ในรายงาน
          #   ใช้ list เมื่อ Root Cause ของเคสนั้นมีมากกว่า 1 ประเด็นที่แยกจากกันได้ชัดเจน
          "photos": [{"browser": str, "path": str}]   # path relative to CWD when script runs, or absolute
        }, ...
      ]
    }
    """
    doc = Document()
    doc.add_heading(f"Photo Evidence — {data['feature']}", level=0)
    doc.add_paragraph(f"รวมหลักฐานภาพทุก Test Case ({len(data['cases'])} เคส)")
    # หน้าปกเป็นหน้าของตัวเอง แยกจากเคสแรกเสมอ เพื่อให้ "1 หน้าต่อ 1 Test Case" เป็นจริงทุกเคส
    # รวมถึงเคสแรกด้วย (ถ้าไม่ page break ตรงนี้ เคสแรกจะไปแชร์หน้ากับหน้าปก ต่างจากเคสอื่นๆ ที่เริ่มหน้าใหม่เสมอ)
    doc.add_page_break()

    for i, c in enumerate(data["cases"]):
        if i > 0:
            doc.add_page_break()
        color = STATUS_COLORS.get(c["status"], RGBColor(0, 0, 0))

        h = doc.add_heading(f"{c['tc_id']} — {c.get('scenario', '')}", level=1)

        meta = doc.add_paragraph()
        status_run = meta.add_run(f"Status: {c['status']}   |   Priority: {c.get('priority', '-')}")
        status_run.bold = True
        status_run.font.color.rgb = color

        doc.add_paragraph(f"คำอธิบายผล: {c.get('description', '-')}")
        _add_issue_points(doc, c.get("issue_point", "-"))

        photos = c.get("photos", [])
        if not photos:
            no_photo = doc.add_paragraph()
            no_photo.add_run("ไม่มีภาพ — ยังไม่ได้ทดสอบ" if c["status"] == "not start" else "ไม่มีภาพแนบ").italic = True
        else:
            for photo in photos:
                if not os.path.exists(photo["path"]):
                    missing = doc.add_paragraph()
                    missing.add_run(f"[ไม่พบไฟล์ภาพที่ระบุ: {photo['path']}]").italic = True
                    continue
                cap = doc.add_paragraph()
                cap.add_run(f"Browser: {photo['browser']}").italic = True
                doc.add_picture(photo["path"], width=Inches(5.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(out_path)
    return out_path


def main():
    p = argparse.ArgumentParser()
    sub = p.add_subparsers(dest="cmd", required=True)

    p_sum = sub.add_parser("summary")
    p_sum.add_argument("--data", required=True)
    p_sum.add_argument("--out", required=True)

    p_photo = sub.add_parser("photo-evidence")
    p_photo.add_argument("--data", required=True)
    p_photo.add_argument("--out", required=True)

    args = p.parse_args()

    with open(args.data, encoding="utf-8") as f:
        data = json.load(f)

    if args.cmd == "summary":
        out = build_summary(data, args.out)
    else:
        out = build_photo_evidence(data, args.out)

    print(json.dumps({"written": out}, ensure_ascii=False))


if __name__ == "__main__":
    main()
