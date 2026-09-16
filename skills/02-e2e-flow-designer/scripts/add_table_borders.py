#!/usr/bin/env python3
"""
จัดตารางในไฟล์ .docx ที่ pandoc สร้างไว้ให้อ่านง่ายขึ้น 2 เรื่อง:

1. เติมกรอบตาราง (grid border ทุกเส้น) — ค่า default ของ pandoc มีแค่เส้นใต้ header
   เส้นเดียว ทำให้ตารางที่มีหลายคอลัมน์ (เช่น สารบัญแหล่งข้อมูล 7 คอลัมน์) อ่านแยก
   แถว/คอลัมน์ยาก
2. ห้ามแถวตารางแตกข้ามหน้า (row cantSplit) + ให้แถว header ขึ้นซ้ำทุกหน้าถ้าตารางยาว
   เกิน 1 หน้า (tblHeader) — ค่า default ของ Word จะปล่อยให้แถวหนึ่งแหว่งไปโผล่อีกหน้า
   ถ้าเนื้อหาในแถวยาว (เช่น คอลัมน์ "หมายเหตุ" ที่ตัวอักษรเยอะ) ทำให้ข้อมูลของแหล่งเดียวกัน
   ถูกตัดกลางแถวไปอยู่คนละหน้า อ่านสับสน

สคริปต์นี้แก้ไฟล์ .docx ที่ pandoc export ไว้แล้วเท่านั้น

ใช้งาน:
    python3 add_table_borders.py "<path .docx>"

แก้ไฟล์ในที่เดิม (in-place) — ควรรันหลัง pandoc export เสร็จแล้วเท่านั้น ห้ามรันกับไฟล์ .docx
ที่เป็นไฟล์หลัก (source of truth ยังคงเป็น .md เท่านั้น ไฟล์นี้แก้แค่สำเนา .docx เพื่อความสวยงาม)
"""
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_table_borders(table, sz="4", color="000000"):
    tbl = table._tbl
    tblPr = tbl.tblPr
    # ลบ tblBorders เดิมถ้ามี กันซ้อนกัน
    existing = tblPr.find(qn("w:tblBorders"))
    if existing is not None:
        tblPr.remove(existing)

    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)  # หน่วย 1/8 pt — "4" = 0.5pt เส้นบาง อ่านง่ายไม่รก
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def set_rows_no_split(table):
    for i, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        # กันแถวนี้แตกข้ามหน้า
        if trPr.find(qn("w:cantSplit")) is None:
            trPr.append(OxmlElement("w:cantSplit"))
        # แถวแรก (header) ให้ขึ้นซ้ำทุกหน้าถ้าตารางยาวเกิน 1 หน้า
        if i == 0 and trPr.find(qn("w:tblHeader")) is None:
            trPr.append(OxmlElement("w:tblHeader"))


def main():
    if len(sys.argv) != 2:
        print("usage: python3 add_table_borders.py <path .docx>", file=sys.stderr)
        sys.exit(1)

    path = sys.argv[1]
    doc = Document(path)
    if not doc.tables:
        print("ไม่พบตารางในไฟล์นี้ — ไม่มีอะไรให้แก้")
        return

    for table in doc.tables:
        set_table_borders(table)
        set_rows_no_split(table)

    doc.save(path)
    print(f"เติมกรอบ + กันแถวแตกข้ามหน้าให้ {len(doc.tables)} ตารางใน {path} เรียบร้อย")


if __name__ == "__main__":
    main()
