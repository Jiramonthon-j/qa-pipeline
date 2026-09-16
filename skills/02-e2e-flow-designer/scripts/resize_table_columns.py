#!/usr/bin/env python3
"""
Mockup: จัดความกว้างคอลัมน์ตารางใน .docx ที่ pandoc สร้างไว้ให้อ่านง่ายขึ้น

ปัญหาเดิม: pandoc หารความกว้างคอลัมน์เท่า ๆ กันทุกคอลัมน์เสมอ ไม่ว่าคอลัมน์นั้นจะเป็น ID สั้น ๆ
หรือรายละเอียดยาวเป็นย่อหน้าก็ตาม ทำให้คอลัมน์ข้อความยาว (เช่น "รายละเอียด", "หมายเหตุ",
"ข้อเสนอแนะ/Mitigation") แคบเกินไปจนคำถูกตัดกลางคำ อ่านทีละ 1-2 คำต่อบรรทัด

วิธีแก้: คำนวณความกว้างคอลัมน์ตามเนื้อหาจริงในตาราง (ไม่ hardcode ตาม header เพราะ Skill 01/02
มีหลายตารางโครงสร้างต่างกัน — ต้องเป็น heuristic ทั่วไปที่ใช้ได้กับทุกตาราง) แล้วตั้งความกว้าง
แบบตายตัว (fixed layout) แทนการปล่อยให้ Word/pandoc หารเท่ากัน

ใช้งาน:
    python3 resize_table_columns.py "<path .docx>"

แก้ไฟล์ .docx ในที่เดิม (in-place) — ควรรันหลัง pandoc export และก่อน add_table_borders.py
"""
import sys

from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn

USABLE_WIDTH_IN = 6.5  # Letter, margin 1" ทั้งสองข้าง — ตั้ง section ให้ตรงกันเสมอ (ดู _ensure_page_setup)
MIN_COL_WIDTH_IN = 0.65
MAX_COL_WIDTH_IN = 3.2


def _ensure_page_setup(doc):
    """pandoc ไม่ได้ระบุ page size/margin ไว้ในไฟล์เสมอไป (ปล่อยให้ผู้เปิดใช้ default ของตัวเอง)
    ทำให้คำนวณความกว้างคอลัมน์จากสมมติฐาน 6.5" คลาดเคลื่อนได้ถ้าเปิดคนละโปรแกรม — ล็อกให้ชัดเจนไปเลย
    เป็น Letter portrait margin 1" ทุกด้าน เพื่อให้ความกว้างคอลัมน์ที่คำนวณไว้ตรงกับพื้นที่จริงเสมอ"""
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)


def _cell_text(cell):
    return " ".join(p.text for p in cell.paragraphs).strip()


def _col_weights(table):
    ncols = len(table.columns)
    header_lens = [len(_cell_text(c)) for c in table.rows[0].cells[:ncols]]
    data_lens_by_col = [[] for _ in range(ncols)]
    for row in table.rows[1:]:
        for i, cell in enumerate(row.cells[:ncols]):
            data_lens_by_col[i].append(len(_cell_text(cell)))
    weights = []
    for i in range(ncols):
        avg_data = sum(data_lens_by_col[i]) / len(data_lens_by_col[i]) if data_lens_by_col[i] else 0
        weights.append(max(header_lens[i], avg_data, 1))
    return weights


def _water_fill(weights, total, min_w, max_w):
    n = len(weights)
    widths = [min_w] * n
    remaining = total - min_w * n
    if remaining <= 0:
        return widths
    active = set(range(n))
    while remaining > 1e-6 and active:
        total_w = sum(weights[i] for i in active)
        if total_w <= 0:
            share = remaining / len(active)
            for i in active:
                widths[i] += share
            remaining = 0
            break
        capped = []
        alloc = {}
        for i in active:
            add = remaining * (weights[i] / total_w)
            if widths[i] + add >= max_w:
                alloc[i] = max_w - widths[i]
                capped.append(i)
            else:
                alloc[i] = add
        for i in active:
            widths[i] += alloc[i]
        remaining -= sum(alloc.values())
        if not capped:
            break
        active -= set(capped)
    return widths


def _set_cell_width(cell, width):
    cell.width = width
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = tcPr.makeelement(qn("w:tcW"), {})
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(width.twips))
    tcW.set(qn("w:type"), "dxa")


def _apply_widths(table, widths_in):
    # จุดสำคัญของ python-docx: ต้องตั้งทั้ง table.columns[i].width (กำหนด w:tblGrid) และ
    # cell.width ของทุก cell (กำหนด w:tcW) คู่กันเสมอ — ตั้งแค่อย่างใดอย่างหนึ่งจะไม่มีผลจริง
    table.autofit = False
    table.allow_autofit = False
    widths = [Inches(w) for w in widths_in]
    for i, w in enumerate(widths):
        if i < len(table.columns):
            table.columns[i].width = w
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            _set_cell_width(cell, w)


def resize_tables(doc):
    n = 0
    for table in doc.tables:
        if not table.rows or not table.columns:
            continue
        weights = _col_weights(table)
        widths_in = _water_fill(weights, USABLE_WIDTH_IN, MIN_COL_WIDTH_IN, MAX_COL_WIDTH_IN)
        _apply_widths(table, widths_in)
        n += 1
    return n


def main():
    if len(sys.argv) != 2:
        print("usage: python3 resize_table_columns.py <path .docx>", file=sys.stderr)
        sys.exit(1)
    path = sys.argv[1]
    doc = Document(path)
    _ensure_page_setup(doc)
    n = resize_tables(doc)
    doc.save(path)
    print(f"ปรับความกว้างคอลัมน์ให้ {n} ตารางใน {path} เรียบร้อย")


if __name__ == "__main__":
    main()
